import os
from PIL import Image as PILImage
from docx import Document
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# Configuration paths (relative to script location)
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_DIR = os.path.join(BASE_DIR, "ss")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

OUTPUT_DOCX_NAME = "arranged_screenshots.docx"
OUTPUT_PDF_NAME = "arranged_screenshots.pdf"

VALID_EXTENSIONS = ('.png', '.jpg', '.jpeg', '.bmp', '.webp', '.tiff', '.gif')


def get_file_time(path: str) -> float:
    """Get file creation time (st_birthtime on macOS/BSD) or modification time."""
    stat = os.stat(path)
    return getattr(stat, 'st_birthtime', stat.st_mtime)


def get_sorted_images(folder_path: str):
    """Retrieve and sort all valid images by creation/modification time."""
    if not os.path.exists(folder_path):
        os.makedirs(folder_path, exist_ok=True)
        return []

    images = [
        os.path.join(folder_path, f)
        for f in os.listdir(folder_path)
        if f.lower().endswith(VALID_EXTENSIONS) and not f.startswith('.')
    ]
    # Sort by timestamp (chronological order)
    images.sort(key=get_file_time)
    return images


def create_docx(image_files, output_path: str):
    """Generate formatted Word document containing all screenshots."""
    doc = Document()

    # Set page margins to 0.6 inches for generous image display space
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    max_width_inches = 7.0
    max_height_inches = 8.5

    for idx, img_path in enumerate(image_files, 1):
        file_name = os.path.basename(img_path)

        # Caption
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{idx}. {file_name}")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(40, 40, 40)

        # Determine dimensions to maintain aspect ratio within page bounds
        try:
            with PILImage.open(img_path) as img:
                w_px, h_px = img.size
                aspect = w_px / h_px

                # Target width/height in inches
                target_w = max_width_inches
                target_h = target_w / aspect

                if target_h > max_height_inches:
                    target_h = max_height_inches
                    target_w = target_h * aspect

                doc.add_picture(img_path, width=Inches(target_w), height=Inches(target_h))
        except Exception as e:
            # Fallback to fixed width if PIL fails
            doc.add_picture(img_path, width=Inches(max_width_inches))

        doc.add_paragraph()  # spacing between items

    doc.save(output_path)


def create_pdf(image_files, output_path: str):
    """Generate formatted PDF document containing all screenshots."""
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    caption_style = ParagraphStyle(
        'ImageCaption',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor='#222222',
        spaceAfter=6,
        keepWithNext=True
    )

    max_w_pt = letter[0] - 72  # 540 pt
    max_h_pt = letter[1] - 120  # account for caption and margins

    story = []

    for idx, img_path in enumerate(image_files, 1):
        file_name = os.path.basename(img_path)
        caption_text = f"<b>{idx}. {file_name}</b>"
        p = Paragraph(caption_text, caption_style)

        try:
            with PILImage.open(img_path) as img:
                w_px, h_px = img.size
                aspect = w_px / h_px

                target_w = max_w_pt
                target_h = target_w / aspect

                if target_h > max_h_pt:
                    target_h = max_h_pt
                    target_w = target_h * aspect

            rl_img = RLImage(img_path, width=target_w, height=target_h)
            # Group caption and image together
            story.append(KeepTogether([p, rl_img, Spacer(1, 16)]))
        except Exception as e:
            print(f"Warning: Could not process {img_path} for PDF: {e}")
            continue

    doc.build(story)


def main():
    os.makedirs(INPUT_DIR, exist_ok=True)
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    print("========================================")
    print(" Screenshot Arranger (DOCX & PDF Generator)")
    print("========================================")
    print(f"Scanning folder: {INPUT_DIR}")

    image_files = get_sorted_images(INPUT_DIR)

    if not image_files:
        print("\n[!] No screenshots found in the 'ss' folder.")
        print(f"    Please paste your screenshot images into:\n    -> {INPUT_DIR}")
        print("    Then re-run this script.\n")
        return

    print(f"\nFound {len(image_files)} image(s):")
    for idx, f in enumerate(image_files, 1):
        print(f"  {idx}. {os.path.basename(f)}")

    docx_path = os.path.join(OUTPUT_DIR, OUTPUT_DOCX_NAME)
    pdf_path = os.path.join(OUTPUT_DIR, OUTPUT_PDF_NAME)

    print("\n[1/2] Generating DOCX...")
    create_docx(image_files, docx_path)
    print(f"  ✓ Saved: {docx_path}")

    print("[2/2] Generating PDF...")
    create_pdf(image_files, pdf_path)
    print(f"  ✓ Saved: {pdf_path}")

    print("\nAll done! Output files are ready in the 'output' folder.")


if __name__ == "__main__":
    main()
